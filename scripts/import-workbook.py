"""Turn the TimeLog table in "Timesheet- Version 02.xlsm" into SQL for the portal.

    python scripts/import-workbook.py "C:/path/Timesheet- Version 02.xlsm" > timelog-import.sql

Paste the output into the Supabase SQL editor AFTER every person in the
workbook has a portal account whose display name matches their workbook name
(Setup -> Staff). Rows for anyone without an account are skipped and listed
at the end, so it is safe to re-run once they exist: rows already imported
are recognised and not duplicated.

Status mapping (workbook -> portal):
    Invoiced? = Yes  -> approved   (invoices themselves are not migrated)
    Approved By set  -> approved
    otherwise        -> submitted  (shows up in the Approvals queue)
The rate the workbook calculated is frozen on each approved row, so history
keeps its original value even if a project rate changes later.
"""
import datetime as dt
import sys
import warnings

import openpyxl

warnings.filterwarnings("ignore")


def q(v):
    return "null" if v is None or str(v).strip() == "" else "'" + str(v).strip().replace("'", "''") + "'"


def num(v):
    return "null" if not isinstance(v, (int, float)) else repr(round(float(v), 2))


def main(path):
    wb = openpyxl.load_workbook(path, data_only=True, keep_vba=True)
    ws = wb["Timesheet"]
    table = ws.tables["TimeLog"]
    first_col, first_row, last_col, last_row = openpyxl.utils.range_boundaries(table.ref)
    headers = [str(ws.cell(first_row, c).value or "").strip() for c in range(first_col, last_col + 1)]
    col = {h: i for i, h in enumerate(headers)}

    rows = []
    for r in range(first_row + 1, last_row + 1):
        v = [ws.cell(r, c).value for c in range(first_col, last_col + 1)]
        get = lambda h: v[col[h]] if h in col else None
        date, emp, hours = get("Date"), get("Employee"), get("Hours")
        if not isinstance(date, (dt.date, dt.datetime)) or not emp or not isinstance(hours, (int, float)) or hours <= 0:
            continue  # blank or incomplete workbook row
        rows.append((
            date.strftime("%Y-%m-%d"), emp, get("Project No"), get("Project"), get("Work Type"),
            str(get("Chargeable") or "Yes").strip().lower() != "no", hours, get("Task Description") or "(imported from workbook)",
            get("Rate Override"), get("Rate ($/hr)"), get("Approved By"),
            str(get("Invoiced?") or "").strip().lower() == "yes",
        ))

    out = sys.stdout
    out.reconfigure(encoding="utf-8", newline="\n")
    w = out.write
    w(f"-- TimeLog import: {len(rows)} rows from {path.replace(chr(92), '/').split('/')[-1]}, generated {dt.date.today()}\n")
    if not rows:
        w("-- Nothing to import.\n")
        return
    w("begin;\n\ncreate temporary table _timelog (\n  n int, entry_date date, employee text, project_no text, project text, work_type text,\n"
      "  chargeable boolean, hours numeric, description text, rate_override numeric, rate numeric, approved_by text, invoiced boolean\n) on commit drop;\n\n")
    w("insert into _timelog values\n")
    w(",\n".join(
        f"  ({i}, {q(r[0])}, {q(r[1])}, {q(r[2])}, {q(r[3])}, {q(r[4])}, {'true' if r[5] else 'false'}, {num(r[6])}, {q(r[7])}, {num(r[8])}, {num(r[9])}, {q(r[10])}, {'true' if r[11] else 'false'})"
        for i, r in enumerate(rows, 1)))
    w(";\n\n")
    w("""-- Match names to portal records. Work types fall back to 'Original Scope'
-- when blank, because the portal requires one on anything past draft.
create temporary table _matched on commit drop as
select t.*, p.user_id, pr.id as project_id,
       coalesce(wt.id, (select id from public.work_types where name = 'Original Scope')) as work_type_id,
       coalesce(ap.user_id, (select user_id from public.profiles where role = 'admin' order by created_at limit 1)) as approver_id,
       (t.invoiced or t.approved_by is not null) as is_approved
from _timelog t
left join public.profiles   p  on lower(p.display_name) = lower(t.employee)
left join public.projects   pr on pr.project_no = t.project_no
left join public.work_types wt on lower(wt.name) = lower(t.work_type)
left join public.profiles   ap on lower(ap.display_name) = lower(t.approved_by);

with inserted as (
  insert into public.time_entries
    (user_id, entry_date, project_id, work_type_id, chargeable, hours, description, status, approved_by, approved_at)
  select m.user_id, m.entry_date, m.project_id, m.work_type_id, m.chargeable, m.hours, m.description,
         case when m.is_approved then 'approved' else 'submitted' end::public.entry_status,
         case when m.is_approved then m.approver_id end,
         case when m.is_approved then now() end
  from _matched m
  where m.user_id is not null and m.project_id is not null
    and not exists (                       -- safe to re-run
      select 1 from public.time_entries e
      where e.user_id = m.user_id and e.entry_date = m.entry_date and e.project_id = m.project_id
        and e.hours = m.hours and e.description = m.description)
  returning id, user_id, entry_date, project_id, hours, description
)
insert into public.entry_billing (entry_id, rate_override, rate_snapshot)
select i.id, m.rate_override, case when m.is_approved then nullif(m.rate, 0) end
from inserted i
join _matched m on m.user_id = i.user_id and m.entry_date = i.entry_date and m.project_id = i.project_id
               and m.hours = i.hours and m.description = i.description
where m.rate_override is not null or (m.is_approved and nullif(m.rate, 0) is not null)
on conflict (entry_id) do nothing;

-- What could not be imported, and why. Fix (add the person / project) and re-run.
select n as workbook_row, entry_date, employee, project_no, project, hours,
       case when user_id is null then 'no portal account named "' || employee || '"'
            else 'no project numbered "' || coalesce(project_no, '') || '"' end as problem
from _matched where user_id is null or project_id is null
order by n;

commit;
""")


if __name__ == "__main__":
    if len(sys.argv) != 2:
        sys.exit(__doc__)
    main(sys.argv[1])
